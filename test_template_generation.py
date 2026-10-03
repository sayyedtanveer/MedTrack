#!/usr/bin/env python3
"""
Diagnostic script to test Excel template generation
Run this to verify openpyxl and template creation works
"""

import sys
import io

def test_openpyxl_import():
    """Test if openpyxl and its submodules can be imported"""
    print("=" * 60)
    print("TEST 1: Importing openpyxl")
    print("=" * 60)
    
    try:
        import openpyxl
        print(f"✅ openpyxl imported successfully")
        print(f"   Version: {openpyxl.__version__}")
    except ImportError as e:
        print(f"❌ Failed to import openpyxl: {e}")
        return False
    
    try:
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        import openpyxl.comments
        print(f"✅ All required openpyxl submodules imported")
    except ImportError as e:
        print(f"❌ Failed to import openpyxl submodules: {e}")
        return False
    
    return True


def test_workbook_creation():
    """Test creating a multi-sheet workbook"""
    print("\n" + "=" * 60)
    print("TEST 2: Creating Multi-Sheet Workbook")
    print("=" * 60)
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        
        wb = openpyxl.Workbook()
        print(f"✅ Workbook created")
        
        # Test sheet 1
        ws1 = wb.active
        if ws1 is None:
            print(f"❌ Could not access active worksheet")
            return False
        ws1.title = "Field Guide"
        ws1.append(["Header 1", "Header 2", "Header 3"])
        ws1["A1"].font = Font(bold=True, color="FFFFFF")
        ws1["A1"].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        print(f"✅ Sheet 1 'Field Guide' created with styling")
        
        # Test sheet 2
        ws2 = wb.create_sheet(title="Quick Start")
        ws2.append(["Instructions go here"])
        print(f"✅ Sheet 2 'Quick Start' created")
        
        # Test sheet 3
        ws3 = wb.create_sheet(title="Raw Materials")
        ws3.append(["Col1", "Col2", "Col3"])
        print(f"✅ Sheet 3 'Raw Materials' created")
        
        print(f"\n📊 Workbook contains {len(wb.sheetnames)} sheets:")
        for i, name in enumerate(wb.sheetnames, 1):
            print(f"   {i}. {name}")
        
        return wb
        
    except Exception as e:
        print(f"❌ Failed to create workbook: {type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return None


def test_save_to_buffer():
    """Test saving workbook to BytesIO buffer"""
    print("\n" + "=" * 60)
    print("TEST 3: Saving to Memory Buffer")
    print("=" * 60)
    
    try:
        import openpyxl
        
        wb = openpyxl.Workbook()
        ws = wb.active
        if ws:
            ws.title = "Test Sheet"
            ws.append(["Test", "Data"])
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        
        size = buf.getbuffer().nbytes
        print(f"✅ Workbook saved to buffer successfully")
        print(f"   Buffer size: {size:,} bytes")
        
        if size < 1000:
            print(f"⚠️  Warning: Buffer size seems small, might be corrupted")
            return False
        
        return True
        
    except Exception as e:
        print(f"❌ Failed to save to buffer: {type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return False


def test_save_to_file():
    """Test saving workbook to actual file"""
    print("\n" + "=" * 60)
    print("TEST 4: Saving to File")
    print("=" * 60)
    
    filename = "test-template-diagnostic.xlsx"
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        
        wb = openpyxl.Workbook()
        
        # Sheet 1
        ws1 = wb.active
        if ws1:
            ws1.title = "Field Guide"
            ws1.append(["Test Field Guide"])
            ws1["A1"].font = Font(bold=True, size=14)
        
        # Sheet 2
        ws2 = wb.create_sheet(title="Quick Start")
        ws2.append(["Test Instructions"])
        
        # Sheet 3
        ws3 = wb.create_sheet(title="Raw Materials")
        ws3.append(["Column 1", "Column 2", "Column 3"])
        
        wb.save(filename)
        print(f"✅ Workbook saved to file: {filename}")
        
        import os
        if os.path.exists(filename):
            size = os.path.getsize(filename)
            print(f"   File size: {size:,} bytes")
            print(f"\n🎯 SUCCESS! Open '{filename}' and verify 3 tabs exist:")
            print(f"   1. Field Guide")
            print(f"   2. Quick Start")
            print(f"   3. Raw Materials")
            return True
        else:
            print(f"❌ File was not created")
            return False
        
    except Exception as e:
        print(f"❌ Failed to save file: {type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return False


def main():
    """Run all diagnostic tests"""
    print("\n" + "🔬" * 30)
    print("EXCEL TEMPLATE GENERATION DIAGNOSTICS")
    print("🔬" * 30 + "\n")
    
    results = []
    
    # Test 1: Import
    results.append(("Import openpyxl", test_openpyxl_import()))
    
    # Test 2: Create workbook
    if results[0][1]:  # Only if import succeeded
        wb = test_workbook_creation()
        results.append(("Create workbook", wb is not None))
    else:
        results.append(("Create workbook", False))
    
    # Test 3: Save to buffer
    if results[1][1]:  # Only if workbook creation succeeded
        results.append(("Save to buffer", test_save_to_buffer()))
    else:
        results.append(("Save to buffer", False))
    
    # Test 4: Save to file
    if results[1][1]:  # Only if workbook creation succeeded
        results.append(("Save to file", test_save_to_file()))
    else:
        results.append(("Save to file", False))
    
    # Summary
    print("\n" + "=" * 60)
    print("SUMMARY")
    print("=" * 60)
    
    for test_name, passed in results:
        status = "✅ PASS" if passed else "❌ FAIL"
        print(f"{status:8} | {test_name}")
    
    total = len(results)
    passed = sum(1 for _, p in results if p)
    
    print(f"\n{passed}/{total} tests passed")
    
    if passed == total:
        print("\n🎉 All tests passed! Template generation should work in production.")
        return 0
    else:
        print("\n⚠️  Some tests failed. Template generation may not work properly.")
        print("    Check error messages above for details.")
        return 1


if __name__ == "__main__":
    sys.exit(main())
