"""
Standalone script to generate the enhanced material onboarding template
Run this to create the new template with Field Guide tab.
"""
import io
import openpyxl
import openpyxl.comments
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FRIENDLY_COLUMNS = [
    "Item Code", "Material Name", "Category", "Base Unit", "Barcode",
    "Warehouse", "Zone", "Rack / Bin", "Opening Stock", "Minimum Stock",
    "Maximum Stock", "Reorder Level", "Reorder Quantity", "Track by Batch?",
    "Track Expiry?", "Shelf Life", "Track Traceability?", "Quality Check Required?",
    "Preferred Supplier", "Supplier Item Code", "Purchase Unit", "Lead Time (Days)",
    "Minimum Order Quantity", "Length Unit", "Can Be Cut?", "Track Remaining Quantity?",
    "Decimal Places", "Reuse Remaining Material?"
]

def generate_template():
    wb = openpyxl.Workbook()
    
    # ═══════════════════════════════════════════════════════════════
    # 1. FIELD GUIDE SHEET (Primary documentation)
    # ═══════════════════════════════════════════════════════════════
    ws_guide = wb.active
    ws_guide.title = "Field Guide"
    
    # Title
    ws_guide.append(["MedTrack ERP - Raw Material Field Guide"])
    ws_guide["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws_guide["A1"].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    ws_guide.merge_cells('A1:D1')
    
    ws_guide.append([])
    ws_guide.append(["This guide explains each field in the template with examples and best practices."])
    ws_guide["A3"].font = Font(size=10, italic=True, color="64748B")
    ws_guide.merge_cells('A3:D3')
    
    ws_guide.append([])
    
    # Header row for guide
    header_row = ws_guide.max_row + 1
    ws_guide.append(["Field Name", "Required?", "Purpose & Usage", "Example / Notes"])
    for cell in ws_guide[header_row]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Field definitions
    field_guide = [
        ("Item Code", "Optional", 
         "Unique identifier for this material. Use this when UPDATING existing materials. Leave blank for new materials - system will auto-generate.",
         "Example: RM-STEEL-001\n✓ Use when updating\n✗ Leave blank for new"),
        
        ("Material Name", "REQUIRED", 
         "The display name of the material. This appears throughout the system. Must be unique.",
         "Example: Steel Sheet 10mm\nTip: Be descriptive and consistent"),
        
        ("Category", "REQUIRED", 
         "Material category for grouping and reporting. System will auto-create if doesn't exist.",
         "Example: Metals, Plastics, Chemicals\nTip: Use consistent naming"),
        
        ("Base Unit", "REQUIRED", 
         "Primary measurement unit for this material. Cannot be changed after creation.",
         "Example: KG, METER, LITER, PCS\n⚠️ PROTECTED: Choose carefully!"),
        
        ("Barcode", "Optional", 
         "Barcode/SKU for scanner integration. Must be unique if provided.",
         "Example: 1234567890123\nTip: Leave blank if no barcode"),
        
        ("Warehouse", "Optional", 
         "Default storage warehouse location name.",
         "Example: Main Warehouse, Raw Materials Store"),
        
        ("Zone", "Optional", 
         "Specific zone within warehouse for organization.",
         "Example: Zone-A, Zone-B, Chemical Storage"),
        
        ("Rack / Bin", "Optional", 
         "Specific rack or bin location for precise placement.",
         "Example: R1-B2, Shelf-3-A"),
        
        ("Opening Stock", "Optional", 
         "Initial quantity when first creating this material. Only used for NEW materials.",
         "Example: 100\nNote: Ignored for updates"),
        
        ("Minimum Stock", "Optional", 
         "Safety stock level. System alerts when stock falls below this.",
         "Example: 50\nUse: Prevents stockouts"),
        
        ("Maximum Stock", "Optional", 
         "Maximum stock to maintain. Helps with inventory planning.",
         "Example: 500\nUse: Prevents overstocking"),
        
        ("Reorder Level", "Optional", 
         "Trigger point for purchase requisition. When stock hits this, system creates PR.",
         "Example: 100\nTip: Set above Min Stock"),
        
        ("Reorder Quantity", "Optional", 
         "How much to order when reorder level is hit. Suggested quantity for PR.",
         "Example: 200\nTip: Consider MOQ & lead time"),
        
        ("Track by Batch?", "Optional", 
         "YES = Track materials by batch numbers (for traceability, expiry). NO = Track only total quantity.",
         "YES: Medicines, food items, chemicals\nNO: Screws, nuts, generic metals\n⚠️ PROTECTED: Hard to change later!"),
        
        ("Track Expiry?", "Optional", 
         "YES = Material has expiry date (requires batch tracking). NO = No expiry.",
         "YES: Adhesives, chemicals with shelf life\nNO: Metals, plastic sheets\nNote: Auto-enables batch tracking"),
        
        ("Shelf Life", "Optional", 
         "Number of days material is usable after production/receipt. Used with expiry tracking.",
         "Example: 180 (days)\nUse: System calculates expiry date"),
        
        ("Track Traceability?", "Optional", 
         "YES = Track each individual item by serial number. NO = Track by quantity only.",
         "YES: High-value items, warranty items\nNO: Low-value bulk materials\n⚠️ PROTECTED: Choose carefully!"),
        
        ("Quality Check Required?", "Optional", 
         "YES = Requires QC inspection before use. NO = Can use immediately after receipt.",
         "YES: Critical materials, safety items\nNO: Low-risk materials"),
        
        ("Preferred Supplier", "Optional", 
         "Default supplier name for this material. Used in purchase requisitions.",
         "Example: Acme Steel Corp\nTip: Must match existing supplier"),
        
        ("Supplier Item Code", "Optional", 
         "Supplier's product code for this material. Used in POs.",
         "Example: AC-STEEL-500\nTip: Copy from supplier catalog"),
        
        ("Purchase Unit", "Optional", 
         "Unit used when purchasing (if different from base unit). System will handle conversion.",
         "Example: If base=KG, purchase=TON\nLeave blank if same as base"),
        
        ("Lead Time (Days)", "Optional", 
         "Number of days from PO to delivery. Used for MRP planning.",
         "Example: 7\nTip: Get from supplier"),
        
        ("Minimum Order Quantity", "Optional", 
         "Minimum quantity supplier will accept per order.",
         "Example: 100\nUse: Ensures PO meets supplier minimums"),
        
        ("Length Unit", "Optional", 
         "For materials sold by length (pipes, cables). Unit for length measurement.",
         "Example: METER, FOOT\nUse: For cuttable materials only"),
        
        ("Can Be Cut?", "Optional", 
         "YES = Material can be cut into smaller lengths (pipes, fabric). NO = Sold as whole units.",
         "YES: Pipes, cables, fabric, sheets\nNO: Pre-cut items, finished parts"),
        
        ("Track Remaining Quantity?", "Optional", 
         "YES = After cutting, track the unused remainder. NO = Remainder is scrapped.",
         "YES: Expensive materials (minimize waste)\nNO: Cheap materials (not worth tracking)\nRequires 'Can Be Cut?' = YES"),
        
        ("Decimal Places", "Optional", 
         "Number of decimal places for quantity precision. 0 = whole numbers only.",
         "0: PCS (45 pieces)\n2: KG (12.50 kg)\n3: GRAM (0.125 gram)\nDefault: 2"),
        
        ("Reuse Remaining Material?", "Optional", 
         "YES = Unused remainder goes back to stock. NO = Remainder is discarded.",
         "YES: Expensive materials\nNO: Contaminated or low-value scraps\nRequires 'Track Remaining Quantity?' = YES"),
    ]
    
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    for field_name, required, purpose, example in field_guide:
        row_num = ws_guide.max_row + 1
        ws_guide.append([field_name, required, purpose, example])
        
        ws_guide[f"A{row_num}"].font = Font(bold=True, size=10)
        ws_guide[f"A{row_num}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws_guide[f"A{row_num}"].border = thin_border
        
        req_cell = ws_guide[f"B{row_num}"]
        req_cell.alignment = Alignment(horizontal="center", vertical="top")
        req_cell.border = thin_border
        if "REQUIRED" in required:
            req_cell.font = Font(bold=True, color="FFFFFF", size=10)
            req_cell.fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
        elif "PROTECTED" in example:
            req_cell.font = Font(bold=True, color="FFFFFF", size=9)
            req_cell.fill = PatternFill(start_color="EA580C", end_color="EA580C", fill_type="solid")
        else:
            req_cell.font = Font(size=10, color="64748B")
        
        ws_guide[f"C{row_num}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws_guide[f"C{row_num}"].border = thin_border
        ws_guide[f"C{row_num}"].font = Font(size=10)
        
        ws_guide[f"D{row_num}"].alignment = Alignment(vertical="top", wrap_text=True)
        ws_guide[f"D{row_num}"].border = thin_border
        ws_guide[f"D{row_num}"].font = Font(size=9, color="1E40AF")
    
    # Column widths
    ws_guide.column_dimensions['A'].width = 25
    ws_guide.column_dimensions['B'].width = 12
    ws_guide.column_dimensions['C'].width = 55
    ws_guide.column_dimensions['D'].width = 45
    
    # Add legend
    ws_guide.append([])
    legend_row = ws_guide.max_row + 1
    ws_guide.append(["LEGEND:"])
    ws_guide[f"A{legend_row}"].font = Font(bold=True, size=11)
    ws_guide.append(["", "REQUIRED", "Must be provided for all materials"])
    ws_guide[f"B{legend_row + 1}"].fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    ws_guide[f"B{legend_row + 1}"].font = Font(color="FFFFFF", bold=True)
    ws_guide.append(["", "⚠️ PROTECTED", "Cannot be changed easily once set - choose carefully!"])
    ws_guide.append(["", "Optional", "Recommended but not mandatory"])
    
    # ═══════════════════════════════════════════════════════════════
    # 2. QUICK START SHEET
    # ═══════════════════════════════════════════════════════════════
    ws_inst = wb.create_sheet(title="Quick Start")
    
    ws_inst.append(["MedTrack ERP - Quick Start Guide"])
    ws_inst["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws_inst["A1"].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    ws_inst.merge_cells('A1:B1')
    
    ws_inst.append([])
    ws_inst.append(["📋 STEPS TO IMPORT:"])
    ws_inst["A3"].font = Font(bold=True, size=11)
    
    steps = [
        "1. Read the 'Field Guide' tab to understand each field",
        "2. Go to 'Raw Materials' tab",
        "3. Delete the sample row (DUMMY-SAMPLE)",
        "4. Enter your material data (one row per material)",
        "5. Save the file",
        "6. Upload in MedTrack ERP (Inventory → Material Onboarding)",
    ]
    for step in steps:
        ws_inst.append([step])
        ws_inst[f"A{ws_inst.max_row}"].font = Font(size=10)
    
    ws_inst.append([])
    ws_inst.append(["⚠️ IMPORTANT RULES:"])
    ws_inst[f"A{ws_inst.max_row}"].font = Font(bold=True, size=11, color="DC2626")
    
    rules = [
        "✓ Material Name, Category, and Base Unit are REQUIRED",
        "✓ Item Code: Leave blank for NEW materials (auto-generated)",
        "✓ Item Code: Fill in when UPDATING existing materials",
        "✓ Do not change column headers",
        "✓ Use Yes/No dropdowns (not yes/no or Y/N)",
        "✓ Protected fields (Base Unit, Batch Tracking, Traceability) cannot be changed after creation",
        "✗ Do not delete any sheet tabs",
    ]
    for rule in rules:
        ws_inst.append([rule])
    
    ws_inst.column_dimensions['A'].width = 25
    ws_inst.column_dimensions['B'].width = 70
    
    # ═══════════════════════════════════════════════════════════════
    # 3. RAW MATERIALS SHEET (Data entry)
    # ═══════════════════════════════════════════════════════════════
    ws = wb.create_sheet(title="Raw Materials")
    
    # Header
    ws.append(FRIENDLY_COLUMNS)
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Sample row
    ws.append([
        "DUMMY-SAMPLE", "Sample Material (Delete Me)", "Metal", "KG", "",
        "", "", "", "100", "10", "500", "20", "50",
        "No", "No", "", "No", "No",
        "", "", "", "7", "",
        "", "No", "No", "2", "No"
    ])
    
    ws["B2"].comment = openpyxl.comments.Comment(
        "⚠️ DELETE THIS ROW before adding your data!\n\nThis is just a sample to show the format.",
        "MedTrack ERP"
    )
    
    ws.freeze_panes = "A2"
    
    # Column widths
    for col_idx, column in enumerate(FRIENDLY_COLUMNS, 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(column) + 4, 15)
    
    # Data validation
    yes_no_dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    yes_no_dv.error = 'Please select Yes or No from the dropdown'
    yes_no_dv.errorTitle = 'Invalid Value'
    ws.add_data_validation(yes_no_dv)
    
    for col_idx, column in enumerate(FRIENDLY_COLUMNS, 1):
        if column.endswith("?"):
            col_letter = get_column_letter(col_idx)
            yes_no_dv.add(f"{col_letter}2:{col_letter}1048576")
    
    # Save
    wb.save("material-onboarding-template-ENHANCED.xlsx")
    print("✓ Template generated: material-onboarding-template-ENHANCED.xlsx")
    print("  - Field Guide tab: Complete documentation")
    print("  - Quick Start tab: Step-by-step instructions")
    print("  - Raw Materials tab: Data entry sheet")

if __name__ == "__main__":
    generate_template()
