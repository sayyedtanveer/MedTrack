import io
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple, Type

import openpyxl
from pydantic import ValidationError
from sqlalchemy import insert, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.infrastructure.persistence.models.import_models import (
    ImportSessionModel,
    ImportSessionRowModel,
)


class BaseImportService:
    """
    Base service for handling Excel-based bulk master data imports.
    Follows a 2-step process:
    1. Parse and validate -> Store in DB session (preview phase).
    2. Confirm -> Read from DB session -> Bulk insert/update in single transaction.
    """

    def __init__(self, session: AsyncSession, tenant_id: uuid.UUID, user_id: uuid.UUID):
        self.session = session
        self.tenant_id = tenant_id
        self.user_id = user_id
        self.module_name = "BASE" # Override in subclasses

    def get_template_columns(self) -> List[str]:
        """Return the list of columns for the Excel template."""
        raise NotImplementedError

    def get_instructions(self) -> List[Tuple[str, str]]:
        """Return instructions for the template: [(Topic, Detail)]"""
        return [
            ("Mandatory Fields", "Ensure all mandatory columns are filled."),
            ("Duplicate Handling", "Choose Skip, Update, or Fail when confirming import."),
        ]

    async def validate_row(self, row_data: Dict[str, Any]) -> List[Dict[str, str]]:
        """
        Validate a single parsed row.
        Returns a list of error dicts: [{"column": "Email", "value": "xyz", "error": "Invalid format"}]
        Return empty list if valid.
        """
        raise NotImplementedError

    async def bulk_insert(self, valid_rows_data: List[Dict[str, Any]], duplicate_strategy: str) -> None:
        """
        Perform bulk insertion/updating of valid rows.
        """
        raise NotImplementedError

    def generate_template(self) -> bytes:
        """Generate a .xlsx template with Data and Instructions sheets."""
        wb = openpyxl.Workbook()
        
        # Data sheet
        ws_data = wb.active
        ws_data.title = "Data"
        columns = self.get_template_columns()
        ws_data.append(columns)
        
        # Instructions sheet
        ws_instructions = wb.create_sheet(title="Instructions")
        ws_instructions.append(["Topic", "Details"])
        for topic, detail in self.get_instructions():
            ws_instructions.append([topic, detail])
            
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def parse_and_store_preview(self, file_content: bytes) -> uuid.UUID:
        """
        Parse uploaded Excel file, validate rows, and store preview in `import_sessions`.
        Returns the `session_id`.
        """
        wb = openpyxl.load_workbook(filename=io.BytesIO(file_content), data_only=True)
        if "Data" in wb.sheetnames:
            ws = wb["Data"]
        else:
            ws = wb.active
            
        headers = []
        rows_data = []
        
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(cell).strip() if cell is not None else f"Column_{j}" for j, cell in enumerate(row)]
                continue
            
            # Skip entirely empty rows
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
                
            row_dict = {}
            for j, cell in enumerate(row):
                if j < len(headers):
                    val = cell
                    if isinstance(val, str):
                        val = val.strip()
                    row_dict[headers[j]] = val
            rows_data.append((i + 1, row_dict))
            
        # Create session
        import_session = ImportSessionModel(
            tenant_id=self.tenant_id,
            user_id=self.user_id,
            module=self.module_name,
            status="PREVIEW",
            total_rows=len(rows_data),
            valid_rows=0,
            error_rows=0,
        )
        self.session.add(import_session)
        await self.session.flush() # Get session ID
        
        valid_count = 0
        error_count = 0
        
        session_rows = []
        for row_num, row_dict in rows_data:
            errors = await self.validate_row(row_dict)
            is_valid = len(errors) == 0
            
            if is_valid:
                valid_count += 1
            else:
                error_count += 1
                
            session_rows.append(
                ImportSessionRowModel(
                    session_id=import_session.id,
                    row_number=row_num,
                    is_valid=is_valid,
                    data_json=row_dict,
                    errors_json=errors,
                )
            )
            
        import_session.valid_rows = valid_count
        import_session.error_rows = error_count
        
        if session_rows:
            self.session.add_all(session_rows)
            
        await self.session.commit()
        return import_session.id

    async def generate_error_report(self, session_id: uuid.UUID) -> Optional[bytes]:
        """
        Generate an Excel report for rows with errors in the session.
        """
        # Fetch invalid rows
        stmt = select(ImportSessionRowModel).where(
            ImportSessionRowModel.session_id == session_id,
            ImportSessionRowModel.is_valid.is_(False)
        ).order_by(ImportSessionRowModel.row_number)
        
        result = await self.session.execute(stmt)
        invalid_rows = result.scalars().all()
        
        if not invalid_rows:
            return None
            
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Errors"
        
        # Determine headers from the first row's data keys
        data_keys = list(invalid_rows[0].data_json.keys())
        headers = ["Original Row Number"] + data_keys + ["Error Column", "Error Value", "Error Message"]
        ws.append(headers)
        
        for row in invalid_rows:
            base_row_data = [row.row_number] + [row.data_json.get(k, "") for k in data_keys]
            
            # If multiple errors for a row, print them on separate lines
            for err in row.errors_json:
                ws.append(base_row_data + [err.get("column", ""), err.get("value", ""), err.get("error", "")])
                
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def execute_import(self, session_id: uuid.UUID, duplicate_strategy: str) -> None:
        """
        Execute the import for all valid rows in the given session.
        """
        start_time = datetime.now()
        
        # Fetch session
        stmt = select(ImportSessionModel).where(
            ImportSessionModel.id == session_id,
            ImportSessionModel.tenant_id == self.tenant_id,
        )
        result = await self.session.execute(stmt)
        import_session = result.scalar_one_or_none()
        
        if not import_session:
            raise ValueError("Import session not found")
            
        if import_session.status != "PREVIEW":
            raise ValueError(f"Session is in {import_session.status} status, cannot import.")
            
        # Fetch valid rows
        rows_stmt = select(ImportSessionRowModel).where(
            ImportSessionRowModel.session_id == session_id,
            ImportSessionRowModel.is_valid.is_(True)
        ).order_by(ImportSessionRowModel.row_number)
        
        rows_result = await self.session.execute(rows_stmt)
        valid_rows = rows_result.scalars().all()
        
        valid_data = [r.data_json for r in valid_rows]
        
        if valid_data:
            await self.bulk_insert(valid_data, duplicate_strategy)
            
        duration_ms = int((datetime.now() - start_time).total_seconds() * 1000)
        
        import_session.status = "COMPLETED"
        import_session.duplicate_strategy = duplicate_strategy
        import_session.duration_ms = duration_ms
        
        await self.session.commit()
